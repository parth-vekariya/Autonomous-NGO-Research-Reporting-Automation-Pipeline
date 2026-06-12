"""Generate structured Excel spreadsheet from NGO research data."""

from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "S.No",
    "Sector",
    "NGO Name (Full Official)",
    "Short Name / Abbreviation",
    "Year of Establishment",
    "Founding Context",
    "Founder(s)",
    "Founder Background",
    "Mission",
    "Vision",
    "Focus Scope (Narrow/Broad)",
    "Key Program 1",
    "Key Program 2",
    "Key Program 3",
    "Key Program 4",
    "Area of Operation (States/Districts)",
    "Geographic Reach",
    "Website",
    "Facebook",
    "Instagram",
    "X (Twitter)",
    "LinkedIn",
    "YouTube",
    "Impact Stat 1",
    "Impact Stat 2",
    "Impact Stat 3",
    "Registered Office",
    "Phone",
    "Email",
    "Government Grants",
    "CSR Funding",
    "International Donors",
    "FCRA Status",
    "12A / 80G Status",
    "NITI Aayog Darpan ID",
    "Government Tie-ups",
    "Corporate CSR Partners",
    "International Affiliations",
    "Awards & Recognitions",
    "Media Coverage",
    "Credibility Ratings",
]


def _format_programs(programs: list) -> list[str]:
    formatted = []
    for p in programs[:4]:
        formatted.append(f"{p['name']}: {p['description']}")
    while len(formatted) < 4:
        formatted.append("")
    return formatted


def _format_stats(stats: list) -> list[str]:
    formatted = []
    for s in stats[:3]:
        formatted.append(f"{s['metric']}: {s['value']} ({s['period']})")
    while len(formatted) < 3:
        formatted.append("")
    return formatted


def _ngo_to_row(index: int, ngo: dict) -> list:
    programs = _format_programs(ngo.get("programs", []))
    stats = _format_stats(ngo.get("impact_stats", []))

    return [
        index,
        ngo.get("sector", ""),
        ngo.get("full_official_name", ""),
        ngo.get("short_name", ""),
        ngo.get("founding_year", ""),
        ngo.get("founding_context", ""),
        ngo.get("founder_names", ""),
        ngo.get("founder_background", ""),
        ngo.get("mission", ""),
        ngo.get("vision", ""),
        ngo.get("focus_scope", ""),
        programs[0],
        programs[1],
        programs[2],
        programs[3],
        ngo.get("states_districts", ""),
        ngo.get("geographic_reach", ""),
        ngo.get("website", ""),
        ngo.get("facebook", ""),
        ngo.get("instagram", ""),
        ngo.get("twitter_x", ""),
        ngo.get("linkedin", ""),
        ngo.get("youtube", ""),
        stats[0],
        stats[1],
        stats[2],
        ngo.get("registered_office", ""),
        ngo.get("phone", ""),
        ngo.get("email", ""),
        ngo.get("government_grants", ""),
        ngo.get("csr_funding", ""),
        ngo.get("international_donors", ""),
        ngo.get("fcra_status", ""),
        ngo.get("tax_exemption", ""),
        ngo.get("darpan_id", ""),
        ngo.get("government_tieups", ""),
        ngo.get("corporate_csr_partners", ""),
        ngo.get("international_affiliations", ""),
        ngo.get("awards", ""),
        ngo.get("media_coverage", ""),
        ngo.get("credibility_ratings", ""),
    ]


def export_to_excel(ngos: list[dict], output_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "NGO Research Database"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    for idx, ngo in enumerate(ngos, 1):
        row_data = _ngo_to_row(idx, ngo)
        for col, value in enumerate(row_data, 1):
            cell = ws.cell(row=idx + 1, column=col, value=value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border

    for col in range(1, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["I"].width = 35
    ws.column_dimensions["J"].width = 35
    ws.freeze_panes = "A2"

    # Summary sheet
    summary = wb.create_sheet("Sector Summary")
    sectors: dict[str, int] = {}
    for ngo in ngos:
        sector = ngo.get("sector", "Unknown")
        sectors[sector] = sectors.get(sector, 0) + 1

    summary.cell(row=1, column=1, value="Sector").font = Font(bold=True)
    summary.cell(row=1, column=2, value="NGO Count").font = Font(bold=True)
    for i, (sector, count) in enumerate(sorted(sectors.items(), key=lambda x: -x[1]), 2):
        summary.cell(row=i, column=1, value=sector)
        summary.cell(row=i, column=2, value=count)

    summary.column_dimensions["A"].width = 40
    summary.column_dimensions["B"].width = 15

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
